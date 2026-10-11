"""Shared openpyxl style primitives for corporate Excel reports."""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def base_styles():
    thin = Side(border_style="thin", color="CBD5E1")
    return {
        "font_title": Font(name="Arial", size=13, bold=True, color="0F172A"),
        "font_sec": Font(name="Arial", size=10, bold=True, color="1E3A8A"),
        "font_info": Font(name="Arial", size=9, color="334155"),
        "font_header": Font(name="Arial", size=9, bold=True, color="FFFFFF"),
        "font_regular": Font(name="Arial", size=9, color="000000"),
        "font_mono": Font(name="Consolas", size=9, color="000000"),
        "font_mono_bold": Font(name="Consolas", size=9, bold=True, color="000000"),
        "font_bold": Font(name="Arial", size=9, bold=True, color="000000"),
        "font_badge_pass": Font(name="Arial", size=10, bold=True, color="065F46"),
        "font_badge_warn": Font(name="Arial", size=10, bold=True, color="92400E"),
        "font_badge_fail": Font(name="Arial", size=10, bold=True, color="991B1B"),
        "fill_header": PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"),
        "fill_sec": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
        "fill_pass": PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),
        "fill_warn": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        "fill_fail": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "fill_zebra": PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"),
        "fill_highlight": PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid"),
        "border_all": Border(left=thin, right=thin, top=thin, bottom=thin),
        "align_center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "align_left": Alignment(horizontal="left", vertical="center", wrap_text=True),
        "align_right": Alignment(horizontal="right", vertical="center", wrap_text=True),
    }


def styled_cell(ws, row, col, value=None, *, font=None, fill=None, alignment=None, border=None, number_format=None):
    cell = ws.cell(row, col, value)
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if alignment is not None:
        cell.alignment = alignment
    if border is not None:
        cell.border = border
    if number_format is not None:
        cell.number_format = number_format
    return cell
